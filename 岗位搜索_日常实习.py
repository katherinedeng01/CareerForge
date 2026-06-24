from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.styles.colors import Color

wb = Workbook()
ws = wb.active
ws.title = "日常实习岗位"

headers = ["编号","匹配度","岗位名称","公司","城市","薪资","经验要求","发布日期","匹配点","签证/工签","标签","来源平台","链接","备注"]
col_widths = [5, 10, 40, 20, 10, 16, 10, 14, 40, 12, 22, 14, 40, 28]

# 表头样式
hdr_fill = PatternFill("solid", fgColor="4472C4")
hdr_font = Font(bold=True, color="FFFFFF", size=10)
for i, (h, w) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=1, column=i, value=h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 22
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# 颜色
green_fill  = PatternFill("solid", fgColor="E8F5E9")
yellow_fill = PatternFill("solid", fgColor="FFF8E1")
orange_fill = PatternFill("solid", fgColor="FFF3E0")

link_font = Font(color="0563C1", underline="single", size=9)
normal_font = Font(size=9)

rows = [
    # ── 香港 🟢 ──
    [1,"🟢高度匹配","Off-cycle Internship · Global Research","UBS","香港","国际投行","应届/在读","⏳ 滚动招聘（6-12月期）","国投证券行业研究经验完全对口；Excel+Tableau数据分析；英语IELTS 7.5；雅思成绩可投外资","❓","外资·国际投行·Off-cycle","LinkedIn / UBS官网","https://jobs.ubs.com/TGnewUI/Search/home/HomeWithPreLoad?partnerid=25008&siteid=5131&PageType=JobDetails&jobid=329937","6-12个月 Off-cycle，完成后可转正式项目；今年6月可入职"],
    [2,"🟢高度匹配","Off-cycle Internship · Global Markets","UBS","香港","国际投行","应届/在读","⏳ 滚动招聘","金融投资课程背景；数据分析能力；英文流利","❓","外资·国际投行·Off-cycle","LinkedIn / UBS官网","https://hk.linkedin.com/jobs/view/2026-off-cycle-internship-global-markets-hong-kong-at-ubs-4365077147","与Global Research岗位同期，可同步申请"],
    [3,"🟢高度匹配","Research Intern（股票研究/行业研究类）","各家券商/资管（Morningstar/Citadel/Jane Street等）","香港","市场价","应届/在读","实时在招（JobsDB每日更新）","行业研究、报告撰写、Excel财务分析均匹配","❓","外资·金融·滚动招聘","JobsDB","https://hk.jobsdb.com/equity-research-intern-jobs/in-Hong-Kong-SAR","平台聚合106+个在招职位，建议直接筛选"],
    [4,"🟢高度匹配","Market Research Intern（市场研究实习）","多家公司","香港","市场价","应届/在读","实时在招（48个）","行业信息收集与分析能力；结构化报告撰写","❓","外资/港资·多行业","LinkedIn","https://hk.linkedin.com/jobs/market-research-intern-jobs","48个在招，每日新增，点链接直接筛选最新"],
    [5,"🟢高度匹配","Business & Customer Experience Analyst Intern","YATA HK（一田百货）","香港","市场价","在读/应届","2026-05-20 发布（近期）","赛马会客户策略洞察经验直接命中；数据可视化+用户体验分析","❓","港资·零售·兼职友好","LinkedIn","https://hk.linkedin.com/jobs/view/business-customer-experience-analyst-intern-part-time-student-welcome-at-yata-hk-4416132817","Part-time/Student Welcome，时间灵活"],
    [6,"🟢高度匹配","Winter Intern · 研究/金融类（88个）","多家公司","香港","市场价","在读","实时在招（7个近期新增）","金融、研究背景；双语能力","❓","多行业·非暑期","LinkedIn","https://hk.linkedin.com/jobs/winter-intern-jobs","非暑期项目，下半年入职，建议同步关注"],
    [7,"🟢高度匹配","ESG Data Analyst Intern","多家公司","香港","市场价","在读","实时在招（19个，3个近期新增）","毕马威ESG大赛半决赛经历；ESG投资课程背景；Tableau可视化","❓","外资·ESG·可持续金融","LinkedIn","https://hk.linkedin.com/jobs/esg-data-analyst-jobs","ESG背景是差异化优势，强烈推荐"],
    [8,"🟡基本匹配","Investment Intern（投资实习）","多家公司","香港","市场价","在读","实时在招（78个）","金融分析/投资课程；数据分析能力","❓","金融·投资","LinkedIn","https://hk.linkedin.com/jobs/investment-intern-jobs","范围较广，可按公司类型筛选"],
    [9,"🟡基本匹配","Research Intern（综合研究类）","多家公司","香港","市场价","在读","实时在招（147个）","研究报告撰写；行业分析能力","❓","多行业","LinkedIn","https://hk.linkedin.com/jobs/research-intern-jobs","量大，建议配合关键词筛选金融/咨询类"],
    [10,"🟡基本匹配","Data Analyst Intern（数据分析实习）","多家公司","香港","市场价","在读","实时在招（34个）","Excel/Tableau/Python数据分析技能","❓","多行业·数据","LinkedIn","https://hk.linkedin.com/jobs/data-analyst-(internship)-jobs","技能匹配度高，行业可灵活选择"],
    [11,"🟡基本匹配","Strategy Intern（战略实习）","多家公司","香港","市场价","在读","实时在招（8个，1个近期新增）","赛马会战略报告经验；商业案例分析课程","❓","战略·咨询","LinkedIn","https://hk.linkedin.com/jobs/strategy-intern-jobs","数量少但含金量高，建议重点关注"],
    [12,"🟡基本匹配","Internship（综合实习聚合页）","多家公司","香港","市场价","在读","实时在招（425个）","——","❓","多行业","LinkedIn","https://hk.linkedin.com/jobs/internship-jobs","可自行筛选发布日期≤1个月的最新岗位"],
    # ── 深圳 🟢 ──
    [13,"🟢高度匹配","日常实习分析师（科技/行业研究组）","光源资本（精品投行）","深圳","有补贴","在读","实时在招","国投证券研究报告经验+财务数据分析+Excel处理上市公司数据——三项均命中","","精品投行·日常实习·核心岗","Boss直聘","https://m.zhipin.com/zhaopin/f5ff708d0330d93d0XZ42Nm1/","精品投行含金量高；行业研究转投行的过渡好岗"],
    [14,"🟢高度匹配","行业研究实习生","各券商/基金（滚动招聘）","深圳","有补贴","在读","实时在招（Boss直聘聚合）","10+份行业研究报告；锂电池/储能行业深度研究经验；财务分析","","券商/基金·行业研究·日常","Boss直聘","https://www.zhipin.com/zhaopin/bdd49923eaca26b91XB53tS0/","直接对应国投证券背景，强烈推荐"],
    [15,"🟢高度匹配","数据分析师实习生","多家公司","深圳","1-5K/月","在读","实时在招（2154个）","Excel/Python/Tableau/R数据分析全覆盖；可视化经验","","互联网/金融/科技·日常实习","猎聘","https://www.liepin.com/zpshujufenxishixisheng/","数量大，建议按行业（金融/咨询类）筛选"],
    [16,"🟢高度匹配","商业分析实习","多家公司","深圳","有补贴","在读","实时在招","数据整合与量化分析；CEO级别报告撰写（赛马会背景）","","多行业·商业分析","Boss直聘","https://www.zhipin.com/zhaopin/e5cfa136cc24063b1nV-2d60EA~~/","与赛马会经历高度匹配"],
    [17,"🟡基本匹配","战略/投研实习生","各基金公司","深圳","有补贴","在读","实时在招（Gank每日更新）","金融投资+ESG投资课程；数据分析；研究报告","","基金·投研·日常实习","Gank Interview","https://www.gankinterview.cn/campus?tab=internship","可在网站筛选：地区=深圳、类型=日常实习"],
    [18,"🟡基本匹配","咨询实习生（战略分析方向）","贝恩咨询 Bain","深圳/上海","顶级咨询水平","在读","截止2026-08-31（Round 2仍开放）","商业案例分析课程；ESG大赛获奖；结构化报告撰写能力","","MBB咨询·战略","Bain官网","https://www.bain.cn/job.php?act=show&id=32","Round 2 截止8月底，宜尽快投"],
    [19,"🟡基本匹配","实习分析师（各方向）","实习僧平台聚合","深圳","有补贴","在读","实时在招（每日更新）","综合技能均可匹配","","多行业·日常实习","实习僧","https://www.shixiseng.com/","专注实习的垂直平台，可按岗位+深圳筛选"],
    [20,"🟠可以尝试","数据分析实习生（各公司）","深圳多家公司","深圳","110-150元/天","在读（附近高校优先）","实时在招","数据分析技能匹配，行业不限","","互联网·数据","Indeed","https://cn.indeed.com/q-%E6%95%B0%E6%8D%AE%E5%88%86%E6%9E%90%E5%AE%9E%E4%B9%A0%E7%94%9F-%E8%81%8C%E4%BD%8D.html","25个在招，可按城市筛选深圳"],
]

fills = {"🟢高度匹配": green_fill, "🟡基本匹配": yellow_fill, "🟠可以尝试": orange_fill}

for r_idx, row in enumerate(rows, 2):
    fill = fills.get(row[1], PatternFill())
    for c_idx, val in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if c_idx == 13 and val.startswith("http"):  # 链接列
            cell.font = link_font
            cell.hyperlink = val
            cell.value = "点击查看"
        else:
            cell.font = normal_font
    ws.row_dimensions[r_idx].height = 42

fname = "岗位搜索结果_香港深圳日常实习_20260624.xlsx"
wb.save(fname)
print(f"已生成：{fname}")
